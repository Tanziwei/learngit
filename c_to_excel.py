#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
C File Documentation Generator
Reads .c files and extracts function documentation to create an Excel file.
"""

import os
import re
import glob
from typing import List, Dict, Any
try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
except ImportError:
    print("请安装openpyxl库: pip install openpyxl")
    exit(1)


class CFunctionParser:
    """Parser for extracting function information from C files"""
    
    def __init__(self):
        # Regex patterns for parsing
        self.function_id_pattern = r'@note\s+Function\s+ID:\s*([A-Z_0-9]+)'
        self.function_signature_pattern = r'(\w+(?:\s*\*)*)\s+(\w+)\s*\(([^)]*)\)'
        self.param_pattern = r'@param\s+(\w+)\s+(.*?)(?=@param|@return|$)'
        self.return_pattern = r'@return\s+(\w+(?:\s*\*)*)\s+(.*?)(?=@param|@return|$)'
        
    def find_c_files(self, directory: str) -> List[str]:
        """Find all .c files in the given directory"""
        pattern = os.path.join(directory, "*.c")
        return glob.glob(pattern)
    
    def extract_comment_block(self, content: str, start_pos: int) -> str:
        """Extract a complete comment block starting from start_pos"""
        lines = content[start_pos:].split('\n')
        comment_lines = []
        in_comment = False
        
        for line in lines:
            stripped = line.strip()
            if stripped.startswith('/**'):
                in_comment = True
                comment_lines.append(line)
            elif in_comment:
                comment_lines.append(line)
                if stripped.endswith('*/'):
                    break
        
        return '\n'.join(comment_lines)
    
    def parse_function_info(self, comment_block: str, next_lines: str) -> Dict[str, Any]:
        """Parse function information from comment block and following code"""
        info = {
            'design_id': '',
            'interface_name': '',
            'syntax': '',
            'parameters': '',
            'return_value': ''
        }
        
        # Extract Function ID (Design ID)
        function_id_match = re.search(self.function_id_pattern, comment_block, re.IGNORECASE)
        if function_id_match:
            info['design_id'] = function_id_match.group(1)
        
        # Find function signature in the lines following the comment
        lines_after_comment = next_lines.split('\n')[:10]  # Look at next 10 lines
        for line in lines_after_comment:
            line = line.strip()
            if line and not line.startswith('//') and not line.startswith('/*'):
                # Try to match function signature
                func_match = re.search(self.function_signature_pattern, line)
                if func_match:
                    return_type = func_match.group(1).strip()
                    func_name = func_match.group(2).strip()
                    params = func_match.group(3).strip()
                    
                    info['interface_name'] = func_name
                    info['syntax'] = f"{return_type} {func_name}({params});"
                    break
        
        # Parse parameters
        param_info = []
        param_matches = re.finditer(self.param_pattern, comment_block, re.DOTALL)
        for match in param_matches:
            param_name = match.group(1)
            param_desc = match.group(2).strip()
            
            # Parse parameter details
            param_details = self.parse_parameter_details(param_desc)
            param_info.append(f"参数名：{param_name}\n{param_details}")
        
        info['parameters'] = '\n\n'.join(param_info)
        
        # Parse return value
        return_match = re.search(self.return_pattern, comment_block, re.DOTALL)
        if return_match:
            return_type = return_match.group(1).strip()
            return_desc = return_match.group(2).strip()
            return_details = self.parse_return_details(return_desc)
            info['return_value'] = f"类型：{return_type}\n{return_details}"
        
        return info
    
    def parse_parameter_details(self, param_desc: str) -> str:
        """Parse parameter description to extract structured information"""
        details = []
        
        # Look for 输入/输出
        io_match = re.search(r'输入/输出[：:]\s*([^\n]*)', param_desc)
        if io_match:
            details.append(f"输入/输出：{io_match.group(1).strip()}")
        else:
            details.append("输入/输出：")
        
        # Look for 值域范围
        range_match = re.search(r'值域范围[：:]\s*([^\n]*)', param_desc)
        if range_match:
            details.append(f"值域范围：{range_match.group(1).strip()}")
        else:
            details.append("值域范围：")
        
        # Look for 初始值
        init_match = re.search(r'初始值[：:]\s*([^\n]*)', param_desc)
        if init_match:
            details.append(f"初始值：{init_match.group(1).strip()}")
        else:
            details.append("初始值：")
        
        # Look for 说明
        desc_match = re.search(r'说明[：:]\s*([^\n]*)', param_desc)
        if desc_match:
            details.append(f"说明：{desc_match.group(1).strip()}")
        else:
            details.append("说明：")
        
        return '\n'.join(details)
    
    def parse_return_details(self, return_desc: str) -> str:
        """Parse return value description to extract structured information"""
        details = []
        
        # Look for 值域范围
        range_match = re.search(r'值域范围[：:]\s*([^\n]*)', return_desc)
        if range_match:
            details.append(f"值域范围：{range_match.group(1).strip()}")
        else:
            details.append("值域范围：")
        
        # Look for 说明
        desc_match = re.search(r'说明[：:]\s*([^\n]*)', return_desc)
        if desc_match:
            details.append(f"说明：{desc_match.group(1).strip()}")
        else:
            details.append("说明：")
        
        return '\n'.join(details)
    
    def parse_c_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Parse a single C file and extract function information"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    content = f.read()
            except UnicodeDecodeError:
                print(f"警告: 无法读取文件 {file_path}")
                return []
        
        functions = []
        
        # Find all /** comment blocks
        comment_starts = []
        pos = 0
        while True:
            pos = content.find('/**', pos)
            if pos == -1:
                break
            comment_starts.append(pos)
            pos += 3
        
        # Process each comment block
        for start_pos in comment_starts:
            comment_block = self.extract_comment_block(content, start_pos)
            
            # Check if this comment contains a Function ID
            if re.search(self.function_id_pattern, comment_block, re.IGNORECASE):
                # Get the content after the comment block
                comment_end = start_pos + len(comment_block)
                next_content = content[comment_end:comment_end + 500]  # Next 500 chars
                
                func_info = self.parse_function_info(comment_block, next_content)
                if func_info['interface_name']:  # Only add if we found a function
                    functions.append(func_info)
        
        return functions


class ExcelGenerator:
    """Generate Excel file from parsed function information"""
    
    def __init__(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "函数文档"
        
    def setup_headers(self):
        """Setup Excel headers"""
        headers = ["设计ID", "接口名称", "语法", "参数", "返回值"]
        
        for col, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Set column widths
        self.worksheet.column_dimensions['A'].width = 20  # 设计ID
        self.worksheet.column_dimensions['B'].width = 25  # 接口名称
        self.worksheet.column_dimensions['C'].width = 40  # 语法
        self.worksheet.column_dimensions['D'].width = 50  # 参数
        self.worksheet.column_dimensions['E'].width = 30  # 返回值
    
    def add_function_data(self, functions: List[Dict[str, Any]]):
        """Add function data to Excel"""
        for row, func in enumerate(functions, 2):
            self.worksheet.cell(row=row, column=1, value=func['design_id'])
            self.worksheet.cell(row=row, column=2, value=func['interface_name'])
            self.worksheet.cell(row=row, column=3, value=func['syntax'])
            
            # Set parameters cell with wrap text
            param_cell = self.worksheet.cell(row=row, column=4, value=func['parameters'])
            param_cell.alignment = Alignment(wrap_text=True, vertical='top')
            
            # Set return value cell with wrap text
            return_cell = self.worksheet.cell(row=row, column=5, value=func['return_value'])
            return_cell.alignment = Alignment(wrap_text=True, vertical='top')
            
            # Set row height to accommodate wrapped text
            self.worksheet.row_dimensions[row].height = 60
    
    def save(self, filename: str):
        """Save Excel file"""
        self.workbook.save(filename)


def main():
    """Main function"""
    # Get current directory
    current_dir = os.getcwd()
    print(f"正在扫描目录: {current_dir}")
    
    # Initialize parser
    parser = CFunctionParser()
    
    # Find all .c files
    c_files = parser.find_c_files(current_dir)
    
    if not c_files:
        print("未找到任何.c文件")
        return
    
    print(f"找到 {len(c_files)} 个.c文件:")
    for file in c_files:
        print(f"  - {file}")
    
    # Parse all functions
    all_functions = []
    for c_file in c_files:
        print(f"正在解析: {os.path.basename(c_file)}")
        functions = parser.parse_c_file(c_file)
        all_functions.extend(functions)
        print(f"  找到 {len(functions)} 个函数")
    
    if not all_functions:
        print("未找到任何带有Function ID的函数")
        return
    
    # Generate Excel
    print(f"正在生成Excel文件，共 {len(all_functions)} 个函数...")
    generator = ExcelGenerator()
    generator.setup_headers()
    generator.add_function_data(all_functions)
    
    output_file = "function_documentation.xlsx"
    generator.save(output_file)
    
    print(f"Excel文件已生成: {output_file}")
    print("\n生成的函数列表:")
    for i, func in enumerate(all_functions, 1):
        print(f"{i}. {func['design_id']} - {func['interface_name']}")


if __name__ == "__main__":
    main()
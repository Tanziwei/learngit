#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Test script to verify the generated Excel content
"""

import openpyxl

def test_excel_content():
    """Test the generated Excel file content"""
    try:
        workbook = openpyxl.load_workbook('function_documentation.xlsx')
        worksheet = workbook.active
        
        print("Excel文件内容验证:")
        print("=" * 50)
        
        # Check headers
        headers = []
        for col in range(1, 6):
            headers.append(worksheet.cell(row=1, column=col).value)
        print(f"标题行: {headers}")
        
        # Check data rows
        row = 2
        while worksheet.cell(row=row, column=1).value:
            print(f"\n第{row-1}个函数:")
            print(f"  设计ID: {worksheet.cell(row=row, column=1).value}")
            print(f"  接口名称: {worksheet.cell(row=row, column=2).value}")
            print(f"  语法: {worksheet.cell(row=row, column=3).value}")
            print(f"  参数: {worksheet.cell(row=row, column=4).value}")
            print(f"  返回值: {worksheet.cell(row=row, column=5).value}")
            row += 1
            
        print(f"\n总共解析了 {row-2} 个函数")
        
    except Exception as e:
        print(f"读取Excel文件时出错: {e}")

if __name__ == "__main__":
    test_excel_content()